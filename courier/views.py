from django.shortcuts import render, redirect
from django.contrib.auth.decorators import login_required
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.http import HttpResponse, JsonResponse
from datetime import datetime, date, timedelta
import openpyxl
import json
from bson import ObjectId

from .models import Booking, Customer, HeadOfficeReport, Expense, Invoice, Address
from .forms import BookingForm, ExpenseForm, HeadOfficeReportForm

# ---------- Auth ----------
def login_view(request):
    if request.user.is_authenticated:
        return redirect('dashboard')
    if request.method == 'POST':
        username = request.POST.get('username')
        password = request.POST.get('password')
        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)
            return redirect('dashboard')
        else:
            messages.error(request, 'Invalid credentials')
    return render(request, 'login.html')

def logout_view(request):
    logout(request)
    return redirect('home')

def home(request):
    return render(request, 'home.html')

# ---------- Dashboard ----------
@login_required
def dashboard(request):
    bookings = Booking.get_all()
    total_bookings = len(bookings)
    total_cash = sum(1 for b in bookings if b.account_type == 'cash')
    total_credit = sum(1 for b in bookings if b.account_type == 'credit')
    total_amount = sum(b.amount for b in bookings)
    pending = sum(1 for b in bookings if b.status == 'pending')
    in_transit = sum(1 for b in bookings if b.status == 'in_transit')
    delivered = sum(1 for b in bookings if b.status == 'delivered')
    issues = sum(1 for b in bookings if b.status == 'issue')
    recent = bookings[:10]

    context = {
        'recent_bookings': recent,
        'total_bookings': total_bookings,
        'total_cash': total_cash,
        'total_credit': total_credit,
        'total_amount': total_amount,
        'pending': pending,
        'in_transit': in_transit,
        'delivered': delivered,
        'issues': issues,
    }
    return render(request, 'dashboard.html', context)

# ---------- Booking CRUD ----------
@login_required
def add_booking(request):
    if request.method == 'POST':
        form = BookingForm(request.POST, request.FILES)
        if form.is_valid():
            data = form.cleaned_data
            # Build address dicts
            from_addr = {
                'name': data['from_address_name'],
                'address_line': data['from_address_address_line'],
                'city': data['from_address_city'],
                'phone': data['from_address_phone'],
            }
            to_addr = {
                'name': data['to_address_name'],
                'address_line': data['to_address_address_line'],
                'city': data['to_address_city'],
                'phone': data['to_address_phone'],
            }
            booking = Booking(
                date=data['date'],
                customer_name=data['customer_name'] or 'Walk-in',
                from_address=Address(**from_addr),
                to_address=Address(**to_addr),
                weight=data['weight'],
                pieces=data['pieces'],
                tracking_id=data['tracking_id'],
                amount=float(data['amount']),
                amount_type=data['amount_type'],
                paid_amount=float(data['paid_amount'] or 0),
                remark=data['remark'],
                account_type=data['account_type'],
                status=data['status'],
                complaint=data['complaint'],
            )
            if data.get('image'):
                # Save image to media folder manually
                booking.image = f'booking_images/{data["image"].name}'
                # Actually write the file
                with open(f'media/{booking.image}', 'wb') as f:
                    for chunk in data['image'].chunks():
                        f.write(chunk)
            booking.save()
            
            from .models import TrackingHistory
            TrackingHistory(
                tracking_id=booking.tracking_id,
                status=booking.status,
                remark='Booking created',
                updated_by=request.user.username if request.user.is_authenticated else 'System'
            ).save()
            
            messages.success(request, 'Booking added!')
            return redirect('booking_list')
    else:
        form = BookingForm()
    return render(request, 'add_booking.html', {'form': form})

@login_required
def edit_booking(request, pk):
    booking = Booking.get(pk)
    if not booking:
        messages.error(request, 'Booking not found')
        return redirect('booking_list')
    if request.method == 'POST':
        form = BookingForm(request.POST, request.FILES)
        if form.is_valid():
            data = form.cleaned_data
            booking.date = data['date']
            booking.customer_name = data['customer_name'] or 'Walk-in'
            booking.from_address = Address(
                name=data['from_address_name'],
                address_line=data['from_address_address_line'],
                city=data['from_address_city'],
                phone=data['from_address_phone']
            )
            booking.to_address = Address(
                name=data['to_address_name'],
                address_line=data['to_address_address_line'],
                city=data['to_address_city'],
                phone=data['to_address_phone']
            )
            booking.weight = data['weight']
            booking.pieces = data['pieces']
            booking.tracking_id = data['tracking_id']
            booking.amount = float(data['amount'])
            booking.amount_type = data['amount_type']
            booking.paid_amount = float(data['paid_amount'] or 0)
            booking.remark = data['remark']
            booking.account_type = data['account_type']
            booking.status = data['status']
            booking.complaint = data['complaint']
            if data.get('image'):
                booking.image = f'booking_images/{data["image"].name}'
                with open(f'media/{booking.image}', 'wb') as f:
                    for chunk in data['image'].chunks():
                        f.write(chunk)
            old_status = booking.status
            booking.save()
            
            if old_status != booking.status:
                from .models import TrackingHistory
                TrackingHistory(
                    tracking_id=booking.tracking_id,
                    status=booking.status,
                    remark='Status updated via Edit Booking',
                    updated_by=request.user.username if request.user.is_authenticated else 'System'
                ).save()
                
            messages.success(request, 'Booking updated!')
            return redirect('booking_list')
    else:
        # Pre-fill form with existing data
        initial = {
            'date': booking.date,
            'customer_name': booking.customer_name,
            'from_address_name': booking.from_address.name,
            'from_address_address_line': booking.from_address.address_line,
            'from_address_city': booking.from_address.city,
            'from_address_phone': booking.from_address.phone,
            'to_address_name': booking.to_address.name,
            'to_address_address_line': booking.to_address.address_line,
            'to_address_city': booking.to_address.city,
            'to_address_phone': booking.to_address.phone,
            'weight': booking.weight,
            'pieces': booking.pieces,
            'tracking_id': booking.tracking_id,
            'amount': booking.amount,
            'amount_type': booking.amount_type,
            'paid_amount': booking.paid_amount,
            'remark': booking.remark,
            'account_type': booking.account_type,
            'status': booking.status,
            'complaint': booking.complaint,
        }
        form = BookingForm(initial=initial)
    return render(request, 'edit_booking.html', {'form': form, 'booking': booking})

@login_required
def booking_list(request):
    bookings = Booking.get_all()
    return render(request, 'booking_list.html', {'bookings': bookings})

@login_required
def update_status(request, pk):
    booking = Booking.get(pk)
    if booking and request.method == 'POST':
        new_status = request.POST.get('status', booking.status)
        complaint = request.POST.get('complaint', '')
        
        if booking.status != new_status or complaint:
            booking.status = new_status
            booking.complaint = complaint
            booking.save()
            
            from .models import TrackingHistory
            TrackingHistory(
                tracking_id=booking.tracking_id,
                status=new_status,
                remark=complaint or 'Status updated',
                updated_by=request.user.username
            ).save()
            
            messages.success(request, 'Status updated!')
    return redirect('booking_list')

# ---------- Head Office Reports ----------
@login_required
def add_ho_report(request):
    if request.method == 'POST':
        form = HeadOfficeReportForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data
            HeadOfficeReport(**data).save()
            messages.success(request, 'HO Report added!')
            return redirect('ho_report_list')
    else:
        form = HeadOfficeReportForm()
    return render(request, 'add_ho_report.html', {'form': form})

@login_required
def ho_report_list(request):
    reports = HeadOfficeReport.get_all()
    return render(request, 'ho_report_list.html', {'reports': reports})

# ---------- Search / Filter ----------
@login_required
def search_view(request):
    query = request.GET.get('q', '')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    status = request.GET.get('status')
    account_type = request.GET.get('account_type')

    # Simple in-memory filtering (MongoDB query could be more efficient)
    bookings = Booking.get_all()
    if query:
        bookings = [b for b in bookings if
                    query.lower() in b.tracking_id.lower() or
                    query.lower() in b.customer_name.lower() or
                    query.lower() in b.from_address.name.lower() or
                    query.lower() in b.to_address.name.lower()
                    ]
    if date_from:
        try:
            dfrom = datetime.strptime(date_from, '%Y-%m-%d').date()
            bookings = [b for b in bookings if b.date >= dfrom]
        except:
            pass
    if date_to:
        try:
            dto = datetime.strptime(date_to, '%Y-%m-%d').date()
            bookings = [b for b in bookings if b.date <= dto]
        except:
            pass
    if status:
        bookings = [b for b in bookings if b.status == status]
    if account_type:
        bookings = [b for b in bookings if b.account_type == account_type]


    context = {
        'bookings': bookings,
        'status': request.GET.get('status', ''),
        'account_type': request.GET.get('account_type', ''),
    }
    return render(request, 'search.html', context)

# ---------- Invoice ----------
@login_required
def create_invoice(request):
    # Get all unique customer names from bookings
    all_bookings = Booking.get_all()
    customers = sorted(list(set(b.customer_name for b in all_bookings if b.customer_name)))
    
    if request.method == 'POST':
        customer_name = request.POST.get('customer')
        date_from = request.POST.get('date_from')
        date_to = request.POST.get('date_to')
        selected = request.POST.getlist('bookings')
        invoice = Invoice(
            customer_name=customer_name,
            date_from=date_from,
            date_to=date_to,
            booking_ids=selected,
        )
        total = 0
        for bid in selected:
            b = Booking.get(bid)
            if b:
                total += b.amount
        invoice.total_amount = total
        invoice.save()
        messages.success(request, 'Invoice created!')
        return redirect('invoice_detail', pk=str(invoice._id))
    
    return render(request, 'create_invoice.html', {'customers': customers})

@login_required
def invoice_detail(request, pk):
    invoice = Invoice.get(pk)
    if not invoice:
        messages.error(request, 'Invoice not found')
        return redirect('invoice_list')
    # Fetch related bookings
    bookings = []
    for bid in invoice.booking_ids:
        b = Booking.get(bid)
        if b:
            bookings.append(b)
    return render(request, 'invoice_detail.html', {'invoice': invoice, 'bookings': bookings})

@login_required
def invoice_list(request):
    invoices = Invoice.get_all()
    return render(request, 'invoice_list.html', {'invoices': invoices})

# ---------- Export ----------
@login_required
def export_excel(request, model):
    response = HttpResponse(content_type='application/ms-excel')
    response['Content-Disposition'] = f'attachment; filename="{model}.xlsx"'
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active

    if model == 'booking':
        ws.title = 'Bookings'
        ws.append(['Tracking ID', 'Date', 'Customer', 'Weight', 'Pieces', 'Amount', 'Paid', 'Extra', 'Inverse', 'Status', 'Account'])
        for b in Booking.get_all():
            ws.append([b.tracking_id, b.date, b.customer_name, b.weight, b.pieces, b.amount, b.paid_amount, b.extra, b.inverse, b.status, b.account_type])
    elif model == 'horeport':
        ws.title = 'HO Reports'
        ws.append(['Date', 'Total Bookings', 'Total Amount', 'Cash Bookings', 'Credit Bookings'])
        for r in HeadOfficeReport.get_all():
            ws.append([r.date, r.total_booking, r.total_amount, r.total_cash_booking, r.total_credit_booking])
    elif model == 'invoice':
        ws.title = 'Invoices'
        ws.append(['ID', 'Customer', 'From', 'To', 'Total', 'Paid'])
        for inv in Invoice.get_all():
            ws.append([str(inv._id), inv.customer_name, inv.date_from, inv.date_to, inv.total_amount, inv.paid])
    else:
        return HttpResponse('Invalid model')
    wb.save(response)
    return response


@login_required
def booking_detail(request, pk):
    booking = Booking.get(pk)
    if not booking:
        messages.error(request, 'Booking not found')
        return redirect('booking_list')
    return render(request, 'booking_detail.html', {'booking': booking})

# ---------- Import ----------
@login_required
def import_excel(request, model):
    if request.method == 'POST' and request.FILES.get('file'):
        file = request.FILES['file']
        if file.name.endswith('.xlsx'):
            wb = openpyxl.load_workbook(file)
            sheet = wb.active
            if model == 'booking':
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    try:
                        Booking(
                            date=row[0], customer_name=row[1], tracking_id=row[2],
                            weight=float(row[3]), pieces=int(row[4]), amount=float(row[5]),
                            paid_amount=float(row[6]), account_type=row[7], status=row[8],
                            from_address=Address(name=row[9], address_line=row[10], city=row[11], phone=row[12]),
                            to_address=Address(name=row[13], address_line=row[14], city=row[15], phone=row[16])
                        ).save()
                    except Exception as e:
                        messages.error(request, f'Row error: {e}')
                messages.success(request, 'Bookings imported!')
            elif model == 'horeport':
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    try:
                        HeadOfficeReport(
                            date=row[0], total_booking=row[1], total_cash_booking=row[2],
                            total_credit_booking=row[3], total_pieces=row[4], total_cash_pieces=row[5],
                            total_credit_pieces=row[6], total_amount=float(row[7]), total_cash_amount=float(row[8]),
                            total_credit_amount=float(row[9]), total_extra=float(row[10]), total_inverse=float(row[11]),
                            remark=row[12] if len(row)>12 else ''
                        ).save()
                    except Exception as e:
                        messages.error(request, f'Row error: {e}')
                messages.success(request, 'HO Reports imported!')
            else:
                messages.error(request, 'Invalid model')
            return redirect('booking_list')
        else:
            messages.error(request, 'Only .xlsx files allowed')
    return render(request, 'import.html', {'model': model})


#------------Analytics-----------------

@login_required
def analytics(request):
    from datetime import date, timedelta
    from collections import Counter

    # Date range filters (default: last 30 days)
    today = date.today()
    date_from_str = request.GET.get('date_from')
    date_to_str = request.GET.get('date_to')
    if date_from_str:
        date_from = datetime.strptime(date_from_str, '%Y-%m-%d').date()
    else:
        date_from = today - timedelta(days=30)
    if date_to_str:
        date_to = datetime.strptime(date_to_str, '%Y-%m-%d').date()
    else:
        date_to = today

    # Get bookings within date range
    all_bookings = Booking.get_all()
    bookings = [b for b in all_bookings if date_from <= b.date <= date_to]

    # Get expenses within date range
    all_expenses = Expense.get_all()
    expenses = [e for e in all_expenses if date_from <= e.date <= date_to]

    # Summary metrics
    total_bookings = len(bookings)
    total_revenue = sum(b.amount for b in bookings)
    total_paid = sum(b.paid_amount for b in bookings)
    total_extra = sum(b.extra for b in bookings)
    total_inverse = sum(b.inverse for b in bookings)
    total_expenses = sum(e.amount for e in expenses)
    profit = total_revenue - total_expenses

    # Status distribution
    status_counts = {
        'pending': sum(1 for b in bookings if b.status == 'pending'),
        'in_transit': sum(1 for b in bookings if b.status == 'in_transit'),
        'delivered': sum(1 for b in bookings if b.status == 'delivered'),
        'issue': sum(1 for b in bookings if b.status == 'issue'),
    }

    # Account type distribution
    cash_count = sum(1 for b in bookings if b.account_type == 'cash')
    credit_count = sum(1 for b in bookings if b.account_type == 'credit')

    # Top customers by booking count
    customer_counter = Counter(b.customer_name for b in bookings)
    top_customers = customer_counter.most_common(5)  # list of (name, count)

    # Top customers by revenue
    revenue_by_customer = {}
    for b in bookings:
        revenue_by_customer[b.customer_name] = revenue_by_customer.get(b.customer_name, 0) + b.amount
    top_revenue_customers = sorted(revenue_by_customer.items(), key=lambda x: x[1], reverse=True)[:5]

    # Daily booking count for date range
    day_labels = []
    day_counts = []
    current_date = date_from
    while current_date <= date_to:
        day_labels.append(current_date.strftime('%d %b'))
        count = sum(1 for b in bookings if b.date == current_date)
        day_counts.append(count)
        current_date += timedelta(days=1)

    # Amount type distribution
    amount_type_counter = Counter(b.amount_type for b in bookings)
    amount_type_labels = list(amount_type_counter.keys())
    amount_type_data = list(amount_type_counter.values())

    context = {
        'date_from': date_from,
        'date_to': date_to,
        'total_bookings': total_bookings,
        'total_revenue': total_revenue,
        'total_paid': total_paid,
        'total_extra': total_extra,
        'total_inverse': total_inverse,
        'total_expenses': total_expenses,
        'profit': profit,
        'status_counts': status_counts,
        'cash_count': cash_count,
        'credit_count': credit_count,
        'top_customers': top_customers,
        'top_revenue_customers': top_revenue_customers,
        'day_labels': day_labels,
        'day_counts': day_counts,
        'amount_type_labels': amount_type_labels,
        'amount_type_data': amount_type_data,
    }
    return render(request, 'analytics.html', context)

# ---------- Expense ----------
@login_required
def add_expense(request):
    if request.method == 'POST':
        form = ExpenseForm(request.POST)
        if form.is_valid():
            data = form.cleaned_data
            Expense(**data).save()
            messages.success(request, 'Expense added!')
            return redirect('expense_list')
    else:
        form = ExpenseForm()
    return render(request, 'add_expense.html', {'form': form})

@login_required
def expense_list(request):
    expenses = Expense.get_all()
    return render(request, 'expense_list.html', {'expenses': expenses})

# ---------- Backup ----------
@login_required
def backup_data(request):
    data = {
        'bookings': [vars(b) for b in Booking.get_all()],
        'expenses': [vars(e) for e in Expense.get_all()],
        'invoices': [vars(i) for i in Invoice.get_all()],
    }
    response = HttpResponse(json.dumps(data, default=str), content_type='application/json')
    response['Content-Disposition'] = 'attachment; filename="backup.json"'
    return response

# ---------- Profile ----------
@login_required
def profile(request):
    if request.method == 'POST':
        user = request.user
        user.first_name = request.POST.get('first_name', '')
        user.last_name = request.POST.get('last_name', '')
        user.email = request.POST.get('email', '')
        user.save()
        messages.success(request, 'Profile updated!')
    return render(request, 'profile.html')

# ---------- API for invoice creation ----------
@login_required
def api_bookings(request):
    customer_name = request.GET.get('customer')
    date_from = request.GET.get('date_from')
    date_to = request.GET.get('date_to')
    bookings = Booking.get_all()
    if customer_name:
        bookings = [b for b in bookings if b.customer_name.lower() == customer_name.lower()]
    if date_from:
        try:
            dfrom = datetime.strptime(date_from, '%Y-%m-%d').date()
            bookings = [b for b in bookings if b.date >= dfrom]
        except: pass
    if date_to:
        try:
            dto = datetime.strptime(date_to, '%Y-%m-%d').date()
            bookings = [b for b in bookings if b.date <= dto]
        except: pass
    data = [{
        'id': str(b._id),
        'tracking_id': b.tracking_id,
        'date': b.date.isoformat() if hasattr(b.date, 'isoformat') else b.date,
        'amount': b.amount,
        'status': b.status,
    } for b in bookings]
    return JsonResponse(data, safe=False)

# ---------- Public Tracking ----------
def track_package(request):
    tracking_id = request.GET.get('tracking_id', '').strip()
    booking = None
    history = []
    
    if tracking_id:
        bookings = Booking.filter(tracking_id=tracking_id)
        if bookings:
            booking = bookings[0]
            from .models import TrackingHistory
            history = TrackingHistory.get_history(tracking_id)
            
    return render(request, 'track.html', {
        'booking': booking,
        'history': history,
        'tracking_id': tracking_id,
        'searched': bool(tracking_id)
    })